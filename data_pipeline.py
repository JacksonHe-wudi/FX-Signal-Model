"""
FX_Model.xlsx -> FX_Model_Clean.xlsx + clean_csv/*.csv

Layer 1 (L1_*): cleaned raw panels, rows = Friday dates, cols = currencies/series.
Layer 2 (L2_*): model-ready features derived from L1.

Conventions
-----------
- All weekly tables are aligned to a common Friday grid (W-FRI).
- Each source observation is placed on the first Friday >= its own date stamp
  (as-of alignment, no look-ahead), then forward-filled with a limit:
  weekly sources ffill limit = 2 weeks, monthly sources limit = 8 weeks.
- All spot quotes are harmonized to USD/XXX (up = USD stronger); EUR, GBP,
  AUD, NZD are inverted from their market convention.
- Missing data stays NaN (never interpolated); see the Coverage sheet.

Usage:  python data_pipeline.py <path_to_FX_Model.xlsx> [output_dir]
"""
import sys, os, datetime
import numpy as np
import pandas as pd
import openpyxl

ASIA9 = ['CNH', 'IDR', 'INR', 'KRW', 'MYR', 'PHP', 'SGD', 'THB', 'TWD']
WEEKLY_FFILL, MONTHLY_FFILL = 2, 8


# ---------------------------------------------------------------- loaders
def _rows(ws, min_row):
    for row in ws.iter_rows(min_row=min_row, values_only=True):
        yield row


def load_cvts(ws, header_row=3, data_row=4, date_idx=1):
    """Standard Citi Velocity CVTSHIST sheet: header r3, data r4+, date col B."""
    hdr = next(ws.iter_rows(min_row=header_row, max_row=header_row, values_only=True))
    cols = [(i, h.replace(' - CLOSE', '').strip()) for i, h in enumerate(hdr)
            if i != date_idx and isinstance(h, str) and h.strip() and 'Date' not in h]
    recs = []
    for row in _rows(ws, data_row):
        d = row[date_idx]
        if isinstance(d, datetime.datetime):
            recs.append([d] + [row[i] if i < len(row) else None for i, _ in cols])
    df = pd.DataFrame(recs, columns=['date'] + [c for _, c in cols])
    return df.set_index('date').sort_index().apply(pd.to_numeric, errors='coerce')


def load_block_pairs(ws, pairs, labels, start_row):
    """Sheets holding several (date_col, value_cols...) blocks side by side."""
    out = {}
    for (dcol, vcols), label in zip(pairs, labels):
        recs = []
        for row in _rows(ws, start_row):
            d = row[dcol] if dcol < len(row) else None
            if isinstance(d, datetime.datetime):
                vals = [row[v] if v < len(row) else None for v in vcols]
                recs.append([d] + vals)
        names = list(label) if isinstance(label, list) else [label]
        seen = {}
        for k, n in enumerate(names):          # de-duplicate repeated labels
            seen[n] = seen.get(n, 0) + 1
            if seen[n] > 1:
                names[k] = f'{n}_{seen[n]}'
        df = pd.DataFrame(recs, columns=['date'] + names)
        df = df.set_index('date').sort_index().apply(pd.to_numeric, errors='coerce')
        for n in names:
            out[n] = df[n].dropna()
    return out


# ---------------------------------------------------------------- alignment
def to_friday(series_map, ffill_limit):
    """As-of align a dict of Series onto the common Friday grid."""
    lo = min(s.index.min() for s in series_map.values())
    hi = max(s.index.max() for s in series_map.values())
    grid = pd.date_range(lo, hi, freq='W-FRI')
    out = {}
    for name, s in series_map.items():
        s = s[~s.index.duplicated(keep='last')].sort_index()
        # stamp each obs on the first Friday >= obs date, keep last per Friday
        fri = s.index + pd.to_timedelta((4 - s.index.dayofweek) % 7, unit='D')
        snapped = pd.Series(s.values, index=fri).groupby(level=0).last()
        out[name] = snapped.reindex(grid).ffill(limit=ffill_limit)
    return pd.DataFrame(out, index=grid)


def wide_to_friday(df, ffill_limit=WEEKLY_FFILL):
    return to_friday({c: df[c].dropna() for c in df.columns}, ffill_limit)


# ---------------------------------------------------------------- main build
def build(xlsx_path, outdir):
    wb = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
    L1, L2, notes = {}, {}, []

    # ---- spot (30 ccys), harmonized to USD/XXX ----
    ws = wb['Top 30 FX Spot']
    spot = load_cvts(ws)
    ren = {}
    for c in spot.columns:                       # FX.SPOT.EUR.USD.CITI
        parts = c.split('.')
        base, quote = parts[2], parts[3]
        if base != 'USD':                        # EUR.USD etc -> invert
            spot[c] = 1.0 / spot[c]
            ren[c] = base
        else:
            ren[c] = quote
    spot = spot.rename(columns=ren)
    L1['spot_usd_all30'] = wide_to_friday(spot)
    L1['spot_usd_asia9'] = L1['spot_usd_all30'][[c for c in ASIA9 if c in spot.columns]]

    # ---- forward points TN/1W/1M ----
    fwd = load_cvts(wb['Asia FX Forward'])
    for tenor in ['TN', '1W', '1M']:
        cols = {c.split('.')[4]: c for c in fwd.columns if f'.{tenor}.' in c}
        L1[f'fwd_pts_{tenor.lower()}'] = wide_to_friday(fwd[list(cols.values())]
                                                        .rename(columns={v: k for k, v in cols.items()}))

    # ---- carry (1M implied, ann.) + realized vol ----
    cv = load_cvts(wb['Carry to Vol'])
    carry = {c.split('.')[3]: c for c in cv.columns if c.startswith('FX.CARRY')}
    rvol = {c.split('.')[3]: c for c in cv.columns if 'REALISED' in c or 'REALIS' in c}
    L1['carry_1m_ann'] = wide_to_friday(cv[list(carry.values())].rename(columns={v: k for k, v in carry.items()}))
    L1['vol_realized_1m'] = wide_to_friday(cv[list(rvol.values())].rename(columns={v: k for k, v in rvol.items()}))

    # ---- implied vol / risk reversals (keep all tenors) ----
    vr = load_cvts(wb['Vol & RR'])
    for kind in ['25RR', 'ATM']:
        for tenor in ['1W', '2W', '1M']:
            cols = {c.split('.')[3]: c for c in vr.columns if f'.{kind}.{tenor}.' in c}
            if cols:
                key = f"{'rr25' if kind == '25RR' else 'vol_implied_atm'}_{tenor.lower()}"
                L1[key] = wide_to_friday(vr[list(cols.values())].rename(columns={v: k for k, v in cols.items()}))

    # ---- positioning (Citi PI, real money) ----
    pi = load_cvts(wb['FX Flows & Position'])
    rm = {c.split('.')[1].replace('PI_', ''): c for c in pi.columns if c.endswith('PI_RM')}
    zs = {c.split('.')[1].replace('PI_', ''): c for c in pi.columns if 'ZSCORE' in c}
    L1['pi_realmoney'] = wide_to_friday(pi[list(rm.values())].rename(columns={v: k for k, v in rm.items()}))
    L1['pi_rm_flow_z'] = wide_to_friday(pi[list(zs.values())].rename(columns={v: k for k, v in zs.items()}))

    # ---- economic surprise ----
    esi = load_cvts(wb['Economics Surpirse'])
    cols = {c.split('.')[4].replace('SI_', ''): c for c in esi.columns}
    L1['esi'] = wide_to_friday(esi[list(cols.values())].rename(columns={v: k for k, v in cols.items()}))

    # ---- CDS 5Y ----
    cds = load_cvts(wb['CDS'])
    cmap = {'SBIIN': 'INR_SBI', 'KOREA': 'KRW', 'MALAYS': 'MYR', 'INDON': 'IDR', 'PHILIP': 'PHP'}
    ren = {}
    for c in cds.columns:
        name = c.split('.')[2].split('-')[0]
        ren[c] = cmap.get(name, name)
    L1['cds_5y'] = wide_to_friday(cds.rename(columns=ren))

    # ---- Citi terms of trade ----
    tot = load_cvts(wb['ToT'])
    cols = {c.split('.')[3].replace('CTOT_', '').split(' ')[0]: c for c in tot.columns}
    L1['ctot'] = wide_to_friday(tot[list(cols.values())].rename(columns={v: k for k, v in cols.items()}))

    # ---- regime block ----
    ws = wb['Regime']
    reg = load_block_pairs(ws, [(1, [2, 3, 4, 5, 6])], [['DXY', 'MXWO', 'VIX', 'DB_G10_CARRY', 'JPM_EMVXY']], 8)
    pmi = load_block_pairs(ws, [(8, [9])], ['US_ISM_PMI'], 8)
    L1['regime'] = to_friday(reg, WEEKLY_FFILL).join(to_friday(pmi, MONTHLY_FFILL))

    # ---- equity indices (interleaved date/value pairs) ----
    ws = wb['Equity Performance']
    r5 = next(ws.iter_rows(min_row=5, max_row=5, values_only=True))
    tick2ccy = {'SHSZ300': 'CNH', 'NIFTY': 'INR', 'KOSPI': 'KRW', 'TWSE': 'TWD', 'STI': 'SGD',
                'SET': 'THB', 'FBMKLCI': 'MYR', 'JCI': 'IDR', 'PCOMP': 'PHP'}
    pairs, labels = [], []
    for i, v in enumerate(r5):
        if isinstance(v, str) and v.strip():
            pairs.append((i - 1, [i]))
            labels.append(tick2ccy.get(v.split(' ')[0], v.split(' ')[0]))
    eq = load_block_pairs(ws, pairs, labels, 6)
    L1['equity'] = to_friday(eq, WEEKLY_FFILL)

    # ---- current account (monthly, USD bn) ----
    ws = wb['Current Account']
    r5 = next(ws.iter_rows(min_row=5, max_row=5, values_only=True))
    cty2ccy = {'Korea': 'KRW', 'Thailand': 'THB', 'India': 'INR', 'Philippines': 'PHP', 'China': 'CNH',
               'Indonesia': 'IDR', 'Malaysia': 'MYR', 'Singapore': 'SGD', 'Taiwan': 'TWD'}
    names = [cty2ccy.get(v, v) for v in r5 if isinstance(v, str) and v.strip()]
    ca = load_block_pairs(ws, [(1, list(range(2, 2 + len(names))))], [names], 7)
    ca_m = pd.DataFrame(ca)
    L1['current_account_usdbn'] = to_friday(ca, MONTHLY_FFILL)

    # ---- government yields (multi-block, mixed asc/desc) ----
    ws = wb['FX Implied Yield']
    blocks = [(1, [2, 3, 4, 5, 6], 'US', ['1Y', '2Y', '3Y', '5Y', '10Y']),
              (8,  [9, 10, 11, 12],  'KRW', ['1Y', '2Y', '5Y', '10Y']),
              (14, [15, 16, 17, 18], 'THB', ['1Y', '2Y', '5Y', '10Y']),
              (20, [21, 22, 23, 24], 'IDR_GT', ['1Y', '2Y', '5Y', '10Y']),
              (26, [27, 28, 29, 30], 'PHP', ['1Y', '2Y', '5Y', '10Y']),
              (32, [33, 34, 35, 36], 'CNH', ['1Y', '2Y', '5Y', '10Y']),
              (38, [39, 40, 41, 42], 'IDR_BV', ['1Y', '2Y', '5Y', '10Y']),
              (44, [45, 46, 47, 48], 'MYR', ['1Y', '2Y', '5Y', '10Y']),
              (50, [51, 52, 53, 54], 'SGD', ['1Y', '2Y', '5Y', '10Y']),
              (56, [57, 58, 59, 60], 'TWD', ['1Y', '2Y', '5Y', '10Y'])]
    ymap = {}
    for dcol, vcols, cty, tenors in blocks:
        got = load_block_pairs(ws, [(dcol, vcols)], [[f'{cty}_{t}' for t in tenors]], 7)
        ymap.update(got)
    yields = to_friday(ymap, WEEKLY_FFILL)
    for t in ['1Y', '2Y', '5Y', '10Y']:
        cols = [c for c in yields.columns if c.endswith('_' + t)]
        L1[f'govt_yield_{t.lower()}'] = yields[cols].rename(columns=lambda c: c[:-len(t) - 1])

    # ---- commodities (weekly block + KOEISEU monthly) ----
    ws = wb['Commodities']
    r5 = next(ws.iter_rows(min_row=5, max_row=5, values_only=True))
    r6 = next(ws.iter_rows(min_row=6, max_row=6, values_only=True))
    com_cols, labels = [], []
    for i in range(2, 11):
        if r6[i]:
            com_cols.append(i)
            lab = str(r5[i]).strip() if r5[i] else str(r6[i]).strip()
            labels.append(lab.replace(' ', '_').replace('-', '_'))
    com = load_block_pairs(ws, [(1, com_cols)], [labels], 7)
    L1['commodities'] = to_friday(com, WEEKLY_FFILL)
    if len(r6) > 13 and r6[13]:
        ko = load_block_pairs(ws, [(12, [13])], ['KR_SEMI_EXPORT_PX'], 7)
        ko = {k: v for k, v in ko.items() if len(v)}
        if ko:
            L1['kr_semi_export_px'] = to_friday(ko, MONTHLY_FFILL)

    # ---- optional REER sheet (present in updated workbook versions) ----
    if 'REER' in wb.sheetnames:
        try:
            L1['reer'] = wide_to_friday(load_cvts(wb['REER']))
        except Exception as e:
            notes.append(f'REER sheet found but not parsed: {e}')

    # ============================================================ L2 features
    spot9 = L1['spot_usd_asia9']
    logspot = np.log(spot9)
    L2['log_spot'] = logspot
    L2['ret_1w'] = logspot.diff()
    # momentum: 12w return excluding the most recent week
    L2['mom_12w_ex1w'] = logspot.shift(1) - logspot.shift(13)
    L2['skew_26w'] = L2['ret_1w'].rolling(26).skew()
    L2['vol_realized_13w_ann'] = L2['ret_1w'].rolling(13).std() * np.sqrt(52) * 100

    common = [c for c in ASIA9 if c in L1['carry_1m_ann'].columns and c in L1['vol_realized_1m'].columns]
    L2['carry_to_realvol'] = (L1['carry_1m_ann'][common] / L1['vol_realized_1m'][common])

    if 'rr25_1m' in L1:
        rr = L1['rr25_1m']
        L2['rr25_1m_z_52w'] = (rr - rr.rolling(52).mean()) / rr.rolling(52).std()
    L2['esi_chg_4w'] = L1['esi'].diff(4)
    L2['ctot_chg_13w'] = L1['ctot'].diff(13)
    L2['cds_chg_4w'] = L1['cds_5y'].diff(4)
    L2['equity_mom_12w'] = np.log(L1['equity']).shift(1) - np.log(L1['equity']).shift(13)
    L2['ca_yoy_usdbn'] = ca_m.sort_index().diff(12).pipe(
        lambda d: to_friday({c: d[c].dropna() for c in d.columns}, MONTHLY_FFILL))

    # yield curve features vs US
    y2, y10 = L1['govt_yield_2y'], L1['govt_yield_10y']
    slope = (y10 - y2)
    L2['curve_slope_10y2y'] = slope
    L2['slope_diff_vs_us'] = slope.drop(columns=['US'], errors='ignore').sub(slope['US'], axis=0)
    L2['ydiff_2y_vs_us'] = y2.drop(columns=['US'], errors='ignore').sub(y2['US'], axis=0)

    # ============================================================ outputs
    os.makedirs(outdir, exist_ok=True)
    csvdir = os.path.join(outdir, 'clean_csv')
    os.makedirs(csvdir, exist_ok=True)

    coverage = []
    for layer, tables in [('L1', L1), ('L2', L2)]:
        for name, df in tables.items():
            df.index.name = 'date'
            df.to_csv(os.path.join(csvdir, f'{layer}_{name}.csv'))
            for c in df.columns:
                s = df[c].dropna()
                coverage.append([f'{layer}_{name}', c, s.index.min(), s.index.max(),
                                 len(s), round(100 * (1 - len(s) / len(df)), 1)])
    cov = pd.DataFrame(coverage, columns=['Table', 'Series', 'First Observation',
                                          'Last Observation', 'Observations', 'Percent Missing'])

    xlsx_out = os.path.join(outdir, 'FX_Model_Clean.xlsx')
    write_pretty_excel(xlsx_out, L1, L2, cov, notes)
    return xlsx_out, csvdir, cov


# ------------------------------------------------------------ pretty excel
# Only the tables the model actually uses, with full descriptive names.
# (sheet title <= 31 chars, source dict, table key, description, number format)
EXCEL_SPEC = [
    # ---- cleaned market data ----
    ('Spot Rates 30 Currencies', 'L1', 'spot_usd_all30',
     'Weekly spot, foreign currency per 1 USD (rise = USD stronger). Input to the global PCA factors.', '0.0000'),
    ('Spot Rates Asia 9',        'L1', 'spot_usd_asia9',
     'Weekly spot for the nine Asian target currencies, foreign currency per 1 USD.', '0.0000'),
    ('Forward Points 1 Week',    'L1', 'fwd_pts_1w',
     'One-week forward points in pips. Weekly roll instrument for the strategy.', '0.0000'),
    ('Implied Carry 1 Month',    'L1', 'carry_1m_ann',
     'Forward-implied carry versus USD, one month, annualized percent (Citi).', '0.00'),
    ('Realized Volatility 1 Month', 'L1', 'vol_realized_1m',
     'One-month realized volatility, annualized percent. Denominator of the carry-to-volatility ratio.', '0.00'),
    ('Implied Volatility 1 Month',  'L1', 'vol_implied_atm_1m',
     'One-month at-the-money implied volatility, annualized percent. Used for position sizing.', '0.00'),
    ('Risk Reversal 25 Delta 1 Month', 'L1', 'rr25_1m',
     'One-month 25-delta risk reversal (vol points). Options-market sentiment.', '0.000'),
    ('Positioning Real Money',   'L1', 'pi_realmoney',
     'Citi positioning indicator, real money investors.', '0.00'),
    ('Positioning Flow Z Score', 'L1', 'pi_rm_flow_z',
     'Citi real money flow z-score.', '0.00'),
    ('Economic Surprise Index',  'L1', 'esi',
     'Citi Economic Surprise Index by economy. Positive relation to the currency expected.', '0.0'),
    ('Sovereign CDS 5 Year',     'L1', 'cds_5y',
     'Five-year sovereign CDS par spread, basis points (USD, senior, CR14). India proxied by State Bank of India.', '0.0'),
    ('Citi Terms of Trade',      'L1', 'ctot',
     'Citi commodity terms-of-trade index by currency.', '0.00'),
    ('Equity Indices',           'L1', 'equity',
     'Local benchmark equity index level, mapped to its currency.', '#,##0.00'),
    ('Current Account USD Billion', 'L1', 'current_account_usdbn',
     'Current account balance, billions of USD, monthly, forward-filled onto Fridays.', '0.0'),
    ('Government Yield 2 Year',  'L1', 'govt_yield_2y',
     'Two-year government bond yield, percent. Includes the United States.', '0.000'),
    ('Government Yield 5 Year',  'L1', 'govt_yield_5y',
     'Five-year government bond yield, percent. Includes the United States.', '0.000'),
    ('Government Yield 10 Year', 'L1', 'govt_yield_10y',
     'Ten-year government bond yield, percent. Includes the United States.', '0.000'),
    ('Commodities',              'L1', 'commodities',
     'Weekly commodity prices: coal, Brent oil, gold, rice, copper, iron ore, semiconductor equities, LNG.', '#,##0.00'),
    ('Korean Semiconductor Prices', 'L1', 'kr_semi_export_px',
     'Korea semiconductor export price index, monthly, forward-filled onto Fridays.', '0.0'),
    ('Regime Indicators',        'L1', 'regime',
     'Dollar index, MSCI World, VIX, Deutsche Bank G10 carry index, JPMorgan EM FX volatility, US ISM PMI.', '0.00'),
    ('Real Effective Exchange Rate', 'L1', 'reer',
     'Real effective exchange rate index by currency.', '0.00'),
    # ---- model features ----
    ('Log Spot Asia 9',          'L2', 'log_spot',
     'Natural log of spot. Dependent variable of the fair-value (cointegration) layer.', '0.0000'),
    ('Weekly Return',            'L2', 'ret_1w',
     'One-week log return of spot (positive = USD stronger).', '0.0000'),
    ('Momentum 12 Weeks',        'L2', 'mom_12w_ex1w',
     'Twelve-week log return excluding the most recent week. Trend signal.', '0.0000'),
    ('Realized Skewness 26 Weeks', 'L2', 'skew_26w',
     'Skewness of weekly returns over a rolling 26-week window. Mean-reversion signal.', '0.000'),
    ('Carry to Volatility Ratio', 'L2', 'carry_to_realvol',
     'One-month implied carry divided by one-month realized volatility. Risk-adjusted carry signal.', '0.000'),
    ('Risk Reversal Z Score',    'L2', 'rr25_1m_z_52w',
     'One-month risk reversal standardized over a rolling 52-week window. Sentiment signal.', '0.00'),
    ('Surprise Index Change 4 Weeks', 'L2', 'esi_chg_4w',
     'Four-week change in the economic surprise index. Fundamental momentum signal.', '0.0'),
    ('Terms of Trade Change 13 Wks', 'L2', 'ctot_chg_13w',
     'Thirteen-week change in the Citi terms-of-trade index.', '0.00'),
    ('CDS Change 4 Weeks',       'L2', 'cds_chg_4w',
     'Four-week change in the five-year sovereign CDS spread, basis points.', '0.0'),
    ('Equity Momentum 12 Weeks', 'L2', 'equity_mom_12w',
     'Twelve-week log return of the local equity index excluding the most recent week.', '0.0000'),
    ('Current Account Yearly Change', 'L2', 'ca_yoy_usdbn',
     'Year-over-year change in the current account balance, billions of USD.', '0.0'),
    ('Yield Curve Slope',        'L2', 'curve_slope_10y2y',
     'Ten-year minus two-year government yield by country, percentage points.', '0.000'),
    ('Slope Differential vs US', 'L2', 'slope_diff_vs_us',
     'Local yield-curve slope (10Y minus 2Y) minus the US slope, percentage points.', '0.000'),
    ('Yield Differential 2Y vs US', 'L2', 'ydiff_2y_vs_us',
     'Local two-year government yield minus the US two-year yield, percentage points.', '0.000'),
]


def write_pretty_excel(path, L1, L2, cov, notes):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    TITLE_F = Font(name='Arial', size=11, bold=True, color='1F3864')
    HDR_F = Font(name='Arial', size=10, bold=True, color='FFFFFF')
    HDR_FILL = PatternFill('solid', fgColor='1F3864')
    DATA_F = Font(name='Arial', size=10)
    THIN = Border(bottom=Side(style='thin', color='D9D9D9'))
    RAW_TAB, FEAT_TAB, META_TAB = '2E75B6', '548235', '808080'

    wb = Workbook()

    # ---- contents sheet ----
    ws = wb.active
    ws.title = 'Contents'
    ws.sheet_properties.tabColor = META_TAB
    ws['A1'] = 'FX Asia Weekly Model - Cleaned Database'
    ws['A1'].font = Font(name='Arial', size=14, bold=True, color='1F3864')
    conventions = [
        'All tables: rows are Friday dates, columns are currencies or series.',
        'Spot convention: foreign currency per 1 USD (a rise means a stronger USD); EUR, GBP, AUD, NZD inverted from market quotation.',
        'Alignment: each observation is placed on the first Friday on or after its own date stamp - no look-ahead.',
        f'Forward fill limits: weekly sources up to {WEEKLY_FFILL} weeks, monthly sources up to {MONTHLY_FFILL} weeks; gaps beyond that stay blank.',
        'Blank cells mean the data does not exist for that date - nothing is interpolated.',
        'Generated by data_pipeline.py - do not edit by hand.'] + notes
    r = 3
    for c_ in conventions:
        ws.cell(row=r, column=1, value=c_).font = DATA_F
        r += 1
    r += 1
    ws.cell(row=r, column=1, value='Sheet').font = HDR_F
    ws.cell(row=r, column=1).fill = HDR_FILL
    ws.cell(row=r, column=2, value='Description').font = HDR_F
    ws.cell(row=r, column=2).fill = HDR_FILL
    r += 1
    for title, layer, key, desc, _ in EXCEL_SPEC:
        src = L1 if layer == 'L1' else L2
        if key in src:
            ws.cell(row=r, column=1, value=title).font = Font(name='Arial', size=10, bold=True)
            ws.cell(row=r, column=2, value=desc).font = DATA_F
            r += 1
    ws.column_dimensions['A'].width = 34
    ws.column_dimensions['B'].width = 110

    # ---- data sheets ----
    for title, layer, key, desc, numfmt in EXCEL_SPEC:
        src = L1 if layer == 'L1' else L2
        if key not in src:
            continue
        df = src[key].dropna(how='all')
        ws = wb.create_sheet(title[:31])
        ws.sheet_properties.tabColor = RAW_TAB if layer == 'L1' else FEAT_TAB
        ws['A1'] = f'{title} - {desc}'
        ws['A1'].font = TITLE_F
        ws.cell(row=2, column=1, value='Date').font = HDR_F
        ws.cell(row=2, column=1).fill = HDR_FILL
        for j, c in enumerate(df.columns, start=2):
            cell = ws.cell(row=2, column=j, value=str(c))
            cell.font, cell.fill = HDR_F, HDR_FILL
            cell.alignment = Alignment(horizontal='center')
            ws.column_dimensions[get_column_letter(j)].width = 12
        ws.column_dimensions['A'].width = 12
        for i, (dt, row) in enumerate(df.iterrows(), start=3):
            dcell = ws.cell(row=i, column=1, value=dt)
            dcell.number_format, dcell.font, dcell.border = 'yyyy-mm-dd', DATA_F, THIN
            for j, v in enumerate(row.values, start=2):
                if pd.notna(v):
                    cell = ws.cell(row=i, column=j, value=float(v))
                    cell.number_format, cell.font, cell.border = numfmt, DATA_F, THIN
        ws.freeze_panes = 'B3'

    # ---- coverage sheet ----
    ws = wb.create_sheet('Data Coverage')
    ws.sheet_properties.tabColor = META_TAB
    ws['A1'] = 'Data Coverage - first and last observation and share of missing weeks per series'
    ws['A1'].font = TITLE_F
    for j, c in enumerate(cov.columns, start=1):
        cell = ws.cell(row=2, column=j, value=c)
        cell.font, cell.fill = HDR_F, HDR_FILL
        ws.column_dimensions[get_column_letter(j)].width = 22
    for i, (_, row) in enumerate(cov.iterrows(), start=3):
        for j, v in enumerate(row.values, start=1):
            cell = ws.cell(row=i, column=j, value=v)
            cell.font, cell.border = DATA_F, THIN
            if isinstance(v, (pd.Timestamp, datetime.datetime)):
                cell.number_format = 'yyyy-mm-dd'
    ws.freeze_panes = 'A3'

    wb.save(path)


if __name__ == '__main__':
    src = sys.argv[1]
    outdir = sys.argv[2] if len(sys.argv) > 2 else 'clean_output'
    xlsx_out, csvdir, cov = build(src, outdir)
    print(f'written: {xlsx_out}')
    print(f'csv dir: {csvdir}')
    print(cov.to_string(index=False, max_rows=200))
